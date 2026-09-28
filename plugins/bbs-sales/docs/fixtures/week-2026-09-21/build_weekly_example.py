from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

# Fixture copy: writes next to this script so running it never overwrites the real report in
# OneDrive (Customers/State of CA Agencies/Weekly Summaries/2026-09-21 Weekly Quote Report - CA State.xlsx).
import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "2026-09-21 Weekly Quote Report - CA State.xlsx")

KIM, AANYA, PARKER = "Kim Quezon", "Aanya Samarasinghe", "Parker Marshall"
D = lambda m, d: date(2026, m, d)
NASPO = "NASPO 7-25-51-02.03"
NR = "Not recorded"
DOC, BC, FORM, CRM = "Quote document", "BC record only", "Bid form inputs", "CRM estimate only"

# rep, date, agency, facility, buyer, description, quote#, vehicle, value, source, crm#, rcv, sub, folder, notes
Q = [
 (KIM, D(9,21), "CDCR", "Centinela State Prison (CEN)", "Thomas (Ryan) Kemp", "Basketball court - asphalt sealer, crack filler, striping paint", "Q 16155", NASPO, 9615.25, DOC, "RFQ-005130", D(9,15), D(9,21), "CDCR (ALL)/CEN - Centinela State Prison - 887895118/Thomas (Ryan) Kemp/Basket ball court 9.15.2026", "Customer PO 39129"),
 (KIM, D(9,21), "CDCR", "Centinela State Prison (CEN)", "Thomas Hendricks", "Sunglasses, 200 ea (RFQ WH-09-16-26)", "LICW WH-09-16-26", NASPO, 7488.00, DOC, "RFQ-005154", D(9,16), D(9,21), "CDCR (ALL)/CEN - Centinela State Prison - 887895118/Thomas Hendricks/Sunglasses 9.16.2026", "quote-log.md: sent"),
 (KIM, D(9,21), "CDCR", "California Men's Colony (CMC)", "Tiffany Ortiz", "Shoe shine supplies (BSP-09-26-001)", "OBS 4007B response", NR, 430.14, FORM, "RFQ-005153", D(9,16), D(9,21), "CDCR (ALL)/CMC - California Men's Colony 887895111/Tiffany Ortiz/Shoe Shine 9.16.2026", "Value summed from inputs.json line items; no quote PDF in folder"),
 (KIM, D(9,21), "CDCR", "Kern Valley State Prison (KVSP)", "Robert Jones", "Refrigerant recovery machine", "Q 16157", NASPO, 978.02, DOC, "RFQ-005175", D(9,18), D(9,21), "CDCR (ALL)/KVSP - Kern Valley State Prison/Robert Jones/refrigerant recovery machine", "Customer PO 39135"),
 (KIM, D(9,21), "CMD", "Yuba City", "Keren Escobar", "RATS packs, coyote (PO 39080)", "Q 16150", NASPO, 11286.12, DOC, "RFQ-005064", D(9,9), D(9,21), "CMD - 887895101/Keren Escobar/Rats Pack 9.21.2026", "Full cert packet filled"),
 (KIM, D(9,21), "CMD", "CSLO San Luis Obispo", "Jessica Payne", "Whiteboard magnets, 4 pk", "SQ-104947", NR, 74.40, DOC, "RFQ-005164", D(9,17), D(9,21), "CMD - 887895101/Jessica Payne/Metallic Magnets 9.17.2026", ""),
 (KIM, D(9,21), "CHP", "EPD-Collision Investigation Unit (060)", "Cheyenne Staggs", "Bosch OTC Enhanced software renewals x6", "SQ-104958", NR, 8127.00, BC, "RFQ-005189 (likely)", D(9,17), D(9,21), "(not saved - CHP (ALL)/Bosch Software is the expected folder)", "Due COB 9/23 to HQ-MAIT@chp.ca.gov; send not confirmed (Kim's mailbox not visible)"),
 (KIM, D(9,21), "CHP", "(not identified)", "(not identified)", "CRM entry - no quote document found", "(none)", NR, 7000.00, CRM, "RFQ-005186", D(9,17), D(9,21), "(none found)", "Unverified; CRM value is a round estimate"),
 (KIM, D(9,22), "CDCR", "California Men's Colony (CMC)", "Eddie Arcano", "Solenoid valves, tamper screws, wire disconnects", "Q 16159", NASPO, 2364.94, DOC, "RFQ-005191", D(9,22), D(9,22), "CDCR (ALL)/CMC - California Men's Colony 887895111/Eddie Arcano/solenoids 9.22.2026", "Customer PO 39153"),
 (KIM, D(9,22), "CDCR", "California Rehabilitation Center (CRC)", "Helena Valenzuela", "Echo/EGO outdoor power equipment and batteries", "Priced RFQ bid sheet", NASPO, 9253.56, DOC, "RFQ-005201", D(9,22), D(9,23), "CDCR (ALL)/CRC - California Rehabilitation Center/Helena Valenzuela/Echo 9.22.2026", "Prices re-checked after Grainger 9/1 NASPO increase"),
 (KIM, D(9,22), "CHP", "CHP Academy (091)", "Greg Warburton", "Academy supplies (mops and janitorial)", "Q 16160", "Not recorded (no contract # on quote)", 1368.64, DOC, "RFQ-005192", D(9,22), D(9,22), "CHP (ALL) - 887895146/CHP Academy/Greg Warburton/Academy Quote 9.22.2026", "Customer PO 39155"),
 (KIM, D(9,22), "CMD", "(not listed)", "Charles Aguilar", "Work order parts (RFO Non-IT LPA)", "SQ-104967", "NASPO 7-25-51-02", 14068.48, DOC, "RFQ-005182", D(9,21), D(9,22), "CMD - 887895101/Charles Aguilar/Work Order Parts 9.21.2026", "BC header shows $0 and no salesperson"),
 (KIM, D(9,22), "SWRCB", "PPSU Team Z", "TaMara Skondras", "HUANUO monitor stand x4 (26-170-007)", "SQ-104951", "Open market (Non-IT, F&R)", 74.24, DOC, "RFQ-005180", None, D(9,22), "CWB - California Water Boards/TaMara Skondras/Monitor Stand 9.22.2026", "quote-log.md: pending"),
 (KIM, D(9,22), "CAL FIRE", "Atwater Training Center", "Jessica Green", "Gripper straps", "Draft quote (BC SQ-104977)", NR, 1050.00, DOC, "RFQ-005199", D(9,22), D(9,23), "Cal Fire (ALL)/Cal Fire Atwater TC/Jessica Green/Gripper Straps 9.22.2026", "Folder file is named DRAFT; BC SQ-104977 shows $1,000"),
 (KIM, D(9,23), "CDFW", "Blythe", "Kimberly Lopez", "Janitorial and paper supplies (PO 39234)", "Q 16169", NASPO, 1520.82, DOC, "RFQ-005168", D(9,18), D(9,23), "CDFW - (Fish & Wildlife) - 887895157/Kimberly (Kim) Lopez/Final Blythe 9.23.2026", ""),
 (KIM, D(9,23), "CDFW", "South Coast Region 5", "Robert Etcheverry", "Field and office supplies (batteries, Kestrel meters, gloves)", "SQ-104980 (rev. of SQ-104966)", NR, 1992.69, DOC, "RFQ-005169", D(9,18), D(9,23), "CDFW - (Fish & Wildlife) - 887895157/Robert Etcheverry/Office Supplies 9.18.2026", "Final version used; SQ-104966 ($2,335.16) is the earlier draft"),
 (KIM, D(9,23), "SWRCB", "(not listed)", "Kamal Daniel", "Hach reagents (26-456-003)", "SQ-104982 rev1", NR, 566.70, DOC, "RFQ-005206", D(9,23), D(9,25), "CWB - California Water Boards/Kamal Daniel/Hach 9.23.2026", "rev1 dated 9/25; BC header shows $0"),
 (KIM, D(9,23), "CDCR", "Pleasant Valley State Prison (PVSP)", "(not identified)", "BC quote - no folder files", "SQ-104979", NR, 1753.60, BC, "RFQ-005208", D(9,23), D(9,23), "(none found)", ""),
 (KIM, D(9,23), "CDCR", "Substance Abuse Treatment Facility (SATF)", "(not identified)", "BC quote - no folder files", "SQ-104978", NR, 54.64, BC, "", None, None, "(none found)", "Not in CRM"),
 (KIM, D(9,23), "CHP", "Monterey Area (730)", "(not identified)", "BC quote - no folder files", "SQ-104983", NR, 499.80, BC, "", None, None, "(none found)", "Not in CRM"),
 (KIM, D(9,24), "CDCR", "California Men's Colony (CMC)", "Heidi Coyler", "HEK TEK sheet metal screws", "Q 16180", NASPO, 1033.80, DOC, "RFQ-005216", D(9,24), D(9,24), "CDCR (ALL)/CMC - California Men's Colony 887895111/Heidi Coyler/HEK TEK 9.24.2026", "Customer PO 39311"),
 (KIM, D(9,24), "CDCR", "Division of Adult Parole Operations (DAPO)", "Edward Rodriguez", "CCU (original request 9/14)", "SQ-104896", NR, 31084.55, DOC, "", None, None, "CDCR (ALL)/DAPO - Division of Adult Parole Operations/Edward Rodriguez/DAPO CCU 9.14.2026", "Revised quote; no CRM record"),
 (KIM, D(9,24), "CDCR", "Richard J. Donovan (RJD)", "Jose Sagrero", "Housekeeping supplies (mop heads, pads, soap)", "SQ-104993 V2", NR, 17426.18, DOC, "RFQ-005176", D(9,21), D(9,24), "CDCR (ALL)/RJD - Richard J. Donovan Correctional Facility/Jose Sagrero/Housekeeping 9.21.2026", "V2 dated 9/25; BC header shows $0; CRM value $20,000"),
 (KIM, D(9,24), "CMD", "(not listed)", "Linda Magill", "Janitorial supplies (DOL RFQ)", "Q 16189", NASPO, 12980.75, DOC, "RFQ-005219", D(9,24), D(9,25), "CMD - 887895101/Linda Magill/Janitorial Supplies 9.24.2026", "Customer PO 39332; filled packet in FILLED/"),
 (KIM, D(9,25), "CDCR", "Office of Correctional Safety (OCS) HQ", "Cariel Hollmer", "Grainger items", "Q 16191", NASPO, 20932.39, DOC, "RFQ-005225", D(9,24), D(9,25), "CDCR (ALL)/Office of Correctional Safety (OCS)/Cariel Hollmer/Grainger Items 9.24.2026", "Customer PO 39368"),
 (AANYA, D(9,21), "PDC", "Porterville Developmental Center", "Juan Avalos", "DS 1161 request", "16156", "NASPO 7-25-51-02", 1393.52, DOC, "RFQ-005197", D(9,21), D(9,24), "PDC - 887895140/Juan Avalos/09.21", ""),
 (AANYA, D(9,22), "DIR", "(not listed)", "Jacqueline Ko", "RFQ (see folder)", "SQ-104964", NR, 1196.46, DOC, "RFQ-005148", None, D(9,22), "DIR - Industrial Relations/Jacqueline Ko/09.22", ""),
 (AANYA, D(9,22), "DSH", "Coalinga", "Jonathan Castro", "Grip cable parts (\"Gopher\")", "SQ-104968", NR, 211.00, DOC, "", None, None, "DSH (ALL)/DSH - Coalinga - 887895145/Jonathan Castro/09.22 Gopher", "No CRM record"),
 (AANYA, D(9,23), "DSH", "Coalinga", "Creg Clay", "Plumbing RFQ", "SQ-104988", NR, 2127.26, DOC, "RFQ-005224", None, None, "DSH (ALL)/DSH - Coalinga - 887895145/Creg Clay/09.23", "CRM submitted date reads 9/14 (before the record was created) - left blank here"),
 (AANYA, D(9,24), "DSH", "Coalinga", "Creg Clay", "Plumbing supplies", "16179 (duplicate: 16187)", "NASPO 7-25-51-02", 10078.27, DOC, "RFQ-005220", D(9,24), D(9,24), "DSH (ALL)/DSH - Coalinga - 887895145/Creg Clay/09.24", "16179 and 16187 are identical; counted once"),
 (AANYA, D(9,24), "DGS", "New Natural Resources Bldg (715 P St)", "Emanuel Stirbulescu", "Spiral wound gaskets x2 + Crane 4\" OS&Y gate valve", "16190", "NASPO 7-25-51-02", 2700.84, DOC, "RFQ-005222", D(9,23), D(9,24), "DGS - Unknown ProCo/RFQ's/Emanuel Stirbulescu", "Buyer flagged as urgent"),
 (AANYA, D(9,24), "OEP", "Office of Exposition Park Management", "Stacy Hiers", "OEPM 26030 email solicitation", "SQ-104992", NR, 43249.96, DOC, "RFQ-005221", D(9,22), D(9,24), "OEP - Exposition Park/Stacy Hiers/09.23", "BC header shows $0"),
 (AANYA, D(9,25), "DSH", "Metropolitan", "Lydia Vizcaino", "Beds", "SQ-104998", NR, 95829.00, DOC, "RFQ-005231", D(9,22), None, "DSH (ALL)/DSH - Metro - 887895142/Lydia Vizcaino/09.23 Beds", "CRM does NOT show this as submitted - confirm it went out"),
 (AANYA, D(9,25), "DSH", "Metropolitan", "Lydia Vizcaino", "Samsung 43\" TVs x12", "SQ-105002", NR, 7937.88, DOC, "RFQ-005234", D(9,23), D(9,25), "DSH (ALL)/DSH - Metro - 887895142/Lydia Vizcaino/09.25- TVs", "Grainger no-quoted the HDMI cables (quote 2065940807); BC header shows $0"),
]

# rep, agency, buyer, request, received, due, status, crm, next step
OPEN = [
 (KIM, "CHP", "Anthony Bartholomew", "Gun cleaning items - unofficial quote/estimate", D(9,21), D(10,2), "Pricing research started in folder", "", "Finish pricing bands; submit by 10/2"),
 (KIM, "CHP", "Anthony Bartholomew", "EMT opportunity", D(9,23), None, "Grainger quote 2065949219 and sourcing sheet in folder; no BBS quote yet", "", "Build BBS quote from Grainger pricing"),
 (KIM, "CMD", "Patricia Sapien", "Eaton - multiple work orders (RFQ v2)", D(9,23), None, "RFQ in folder; no quote", "", "Price and respond"),
 (KIM, "DOJ", "Adrian DeVoogd", "PO26-0471 (JUS 8933 informal competitive)", D(9,22), None, "RFQ docs in folder; no quote", "RFQ-005194", "Price and respond"),
 (KIM, "CDCR - SATF", "(not identified)", "N95s, 2 pallets - budgetary", D(9,17), None, "BYD discontinued; Cato quoted 3M 1870 at $2,016 cost before shipping (9/22)", "", "Build budgetary quote"),
 (KIM, "CDCR - RJD", "(not identified)", "CRM placeholder ($1)", D(9,21), None, "Open in CRM; request not identified", "RFQ-005174", "Identify request or close record"),
 (KIM, "CMD", "(not identified)", "CRM placeholder ($1)", D(9,23), None, "Open in CRM; request not identified", "RFQ-005207", "Identify request or close record"),
 (KIM, "CAL FIRE - Atwater TC", "(not identified)", "CRM placeholder ($1)", D(9,25), None, "Open in CRM; request not identified", "RFQ-005233", "Identify request or close record"),
 (AANYA, "CalVet - VHC Fresno", "Priscilla Gonzales", "RFQ due 9/23", D(9,21), D(9,23), "PAST DUE - no quote in folder; CRM open at $0", "RFQ-005173", "Confirm whether it was missed or no-bid; record outcome"),
 (AANYA, "CalVet - VHC Yountville", "Anthony Goosby", "Dusters, cart cover, fly swatter", D(9,21), None, "RFQ.xlsx updated 9/25; no quote", "", "Finish pricing and submit"),
 (AANYA, "Caltrans - District 11", "Caty McCarty", "Sourcing links started 9/25", D(9,25), None, "Sourcing in progress", "", "Build quote"),
 (AANYA, "DCA - DOI", "Jessie Pagan", "Peace officer equipment", D(9,21), None, "Sourcing links updated 9/25", "", "Build quote"),
 (AANYA, "DSH - Atascadero", "Isaiah Rodriguez", "#26IR012 ceramics (folder 09.30)", D(9,21), None, "Vendor comparison and cross-reference built 9/25", "", "Finish quote before the 9/30 folder date"),
 (AANYA, "PDC", "Connor Stubbs", "Phone booths", D(9,23), None, "Handed to Aanya by Parker 9/23; no folder files, not in CRM", "", "Create CRM record; source"),
 (AANYA, "DGS - BPO CNRA HQ (715 P St)", "(not identified)", "CRM open record", None, None, "Open in CRM at $0", "RFQ-005223", "Identify request or close record"),
 (AANYA, "DSH - Metropolitan", "(not identified)", "CRM open record", None, None, "Open in CRM at $0", "RFQ-005235", "Identify request or close record"),
 (PARKER, "CAL FIRE - Redding TC", "Brandon Self", "Globe turnouts, 15 sets (NFPA 1970)", D(8,24), None, "Quote 900022 built 9/28 ($58,158 pre-tax), not sent. LN Curtis cost $53,899.99 is over $50K; asked Brandon 9/25 if it must be under $50K", "RFQ-005238", "Wait on Brandon's answer; CRM wrongly shows submitted 9/25"),
 (PARKER, "CHP HQ", "Anthony Bartholomew", "Updated pricing on Aug quotes 104585 (wipes and tourniquets, $144,650) and 104586 (first aid, $25,580.75)", D(9,21), None, "Asked 9/21 whether to combine or split; asked 9/23 who handles procurement (SB vs DVBE option)", "", "Re-price and resend once Tony answers"),
 (PARKER, "CHP HQ", "Anthony Bartholomew", "Avon custom kits", D(9,25), None, "Request sheet has small and large sizes swapped; sent to Avon and Grainger sourcing; asked Tony for due date", "", "Create CRM record and folder; quote when sourcing returns"),
]

# type, issue, where, fix, owner
FLAGS = [
 ("Quote log", "Only 2 of 34 quotes have a quote-log.md row for this week (Hendricks, Skondras)", "Buyer folders", "Add a row per quote: Date / Item / Quote# / Vehicle / Value / Outcome", "Kim, Aanya"),
 ("Quote log", "17 of this week's buyer folders have no quote-log.md at all", "Buyer folders", "Create quote-log.md in each", "Kim, Aanya"),
 ("Misfile", "Copies of DSH Coalinga quotes 16179 / 16187 and the filled plumbing RFQ sit in the Atascadero folder", "DSH - Atascadero/Isaiah Rodriguez/08.28/Claude outputs", "Remove the copies from Atascadero (originals are in Creg Clay/09.24)", "Aanya"),
 ("Duplicate", "Quotes 16179 and 16187 are identical ($10,078.27)", "DSH - Coalinga/Creg Clay/09.24", "Keep one, void the other", "Aanya"),
 ("BC", "Quote headers at $0 while the folder PDF carries a value: SQ-104967, SQ-104982, SQ-104992, SQ-104993, SQ-105002 (also SQ-104981, DSH-M, $0)", "Business Central", "Enter lines or delete unused headers", "Kim, Aanya"),
 ("BC", "SQ-104967 (CMD) has no salesperson", "Business Central", "Set salesperson", "Kim"),
 ("CRM", "RFQ-005224 submitted date is 9/14, before the record was created", "Sales CRM", "Correct to actual submit date", "Aanya"),
 ("CRM", "RFQ-005231 (DSH-M beds, $95,829) not marked submitted though SQ-104998 was built 9/25", "Sales CRM", "Confirm sent; set submitted date", "Aanya"),
 ("CRM", "RFQ-005238 (Cal Fire Redding) marked submitted 9/25 but the quote has not been sent", "Sales CRM", "Clear submitted date", "Parker"),
 ("CRM", "No CRM record: DAPO CCU (SQ-104896), DSH-C Castro (SQ-104968), SATF (SQ-104978), CHP Monterey (SQ-104983), CHP Avon kits, PDC phone booths, CHP gun cleaning, SATF N95s", "Sales CRM", "Create records", "Kim, Aanya, Parker"),
 ("CRM", "Kim's CRM values are round estimates (e.g. RJD $20,000 vs quote $17,426.18; OCS $20,000 vs $20,932.39)", "Sales CRM", "Enter the actual quote total", "Kim"),
 ("Folder", "Quotes with no folder files: CHP Bosch SQ-104958, PVSP SQ-104979, SATF SQ-104978, CHP Monterey SQ-104983, CHP RFQ-005186", "Buyer folders", "Save quote PDF to buyer folder", "Kim"),
 ("Outcome", "BC SQ-105001 (SWRCB, $1,844.84, Accepted 9/25) equals the pre-tax value of TaMara Skondras's 9/4 quote SQ-104846 (26-170-005) - likely a win", "CWB/TaMara Skondras/quote-log.md", "Confirm and set Outcome = won", "Kim"),
 ("Numbering", "CHP_Quote_104585 / 104586 numbers are not BC numbers (BC SQ-104585 is an EMSA $0 draft)", "Parker's CHP quotes", "Re-issue from BC when re-pricing", "Parker"),
]

EXCLUDED = [
 ("Kim", "Bechtel DOE (RFQ-005142), Grainger (RFQ-005196), VA Dallas x3 (SQ-104990), US Navy (SQ-104962), Gallup Indian Medical Center (SQ-104985)", "Not state agencies (federal / prime / private)"),
 ("Kim", "USAF SQ-104999, Thunderbird SQ-105003", "Not state; also built 9/25 but not submitted per CRM"),
 ("Parker", "Bechtel Sentinel x2, Bechtel DOE, UCOR", "Not state agencies"),
 ("Parker", "CHP_Quote_104585 / 104586 sent to Tony 9/21", "August quotes (dated 8/7) re-sent for confirmation - not new this week"),
 ("Parker", "Cal Fire Redding 900022", "Built 9/28, not sent"),
 ("Parker", "LA County Sheriff rapid assault tools (emailed 9/23)", "Local government; credited to Joey in CRM/BC"),
 ("Aanya", "Kristie Phan RFQ-005171 ($873.41)", "Account not identifiable as a state agency - add if it is one"),
 ("Kim", "CHP EMT opportunity", "Folder holds a Grainger quote to BBS, not a BBS quote - listed under Open Requests"),
 ("All", "Joey Magliato's CSU / UC and LADWP / LA County work; body armor (BA-*)", "Out of the CA state book by rule"),
]

# ---------- styles ----------
NAVY, GOLD, LIGHT, ZEBRA = "10112C", "947F32", "EFECE3", "F7F7F7"
F = "Arial"
title_f = Font(name=F, size=16, bold=True, color=NAVY)
sub_f = Font(name=F, size=10, italic=True, color="555555")
h_f = Font(name=F, size=10, bold=True, color="FFFFFF")
b_f = Font(name=F, size=10)
bold_f = Font(name=F, size=10, bold=True)
sec_f = Font(name=F, size=12, bold=True, color=GOLD)
h_fill = PatternFill("solid", fgColor=NAVY)
tot_fill = PatternFill("solid", fgColor=LIGHT)
z_fill = PatternFill("solid", fgColor=ZEBRA)
thin = Side(style="thin", color="D0D0D0")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
top2 = Border(top=Side(style="medium", color=NAVY))
USD = '$#,##0.00;($#,##0.00);"-"'
USD0 = '$#,##0;($#,##0);"-"'
DT = "mm/dd/yyyy"
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="top")

def header(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font, c.fill, c.border = h_f, h_fill, box
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30

def body(c, fmt=None, zebra=False, align=wrap):
    c.font, c.border, c.alignment = b_f, box, align
    if fmt: c.number_format = fmt
    if zebra: c.fill = z_fill

wb = Workbook()

# ---------- Quotes ----------
qs = wb.active; qs.title = "Quotes"
qh = ["#", "Rep", "Quote date", "Agency", "Facility / unit", "Buyer", "Description", "Quote #", "Vehicle",
      "Value, pre-tax", "Value source", "Confidence", "CRM RFQ #", "CRM received", "CRM submitted",
      "Days received to submitted", "Folder (under All RFQ's by Agency)", "Notes"]
header(qs, 1, qh, [5, 18, 11, 10, 26, 20, 38, 22, 22, 14, 17, 11, 17, 12, 12, 12, 48, 44])
conf = {DOC: "High", FORM: "Medium", BC: "Medium", CRM: "Low"}
Q.sort(key=lambda r: ({KIM: 0, AANYA: 1, PARKER: 2}[r[0]], r[1], r[2]))
for i, r in enumerate(Q, 2):
    rep, dt, ag, fac, buyer, desc, qn, veh, val, src, crm, rcv, sub, folder, note = r
    vals = [i - 1, rep, dt, ag, fac, buyer, desc, qn, veh, val, src, conf[src], crm, rcv, sub,
            f'=IF(AND(ISNUMBER(N{i}),ISNUMBER(O{i})),O{i}-N{i},"")', folder, note]
    z = i % 2 == 1
    for j, v in enumerate(vals, 1):
        c = qs.cell(row=i, column=j, value=v)
        fmt = DT if j in (3, 14, 15) else USD if j == 10 else "0" if j == 16 else None
        body(c, fmt, z, center if j in (1, 3, 12, 14, 15, 16) else wrap)
QLAST = len(Q) + 1
tr = QLAST + 1
qs.cell(row=tr, column=2, value="Total").font = bold_f
qs.cell(row=tr, column=3, value=f"=COUNTA(B2:B{QLAST})&\" quotes\"").font = bold_f
c = qs.cell(row=tr, column=10, value=f"=SUM(J2:J{QLAST})"); c.font = bold_f; c.number_format = USD
for j in range(1, len(qh) + 1):
    qs.cell(row=tr, column=j).fill = tot_fill; qs.cell(row=tr, column=j).border = top2
qs.freeze_panes = "C2"
qs.auto_filter.ref = f"A1:R{QLAST}"
qs["J1"].comment = Comment("Pre-tax subtotal read from the quote PDF where one exists; otherwise the BC quote total, bid-form inputs, or the CRM estimate (see Value source).", "Claude")
qs["P1"].comment = Comment("CRM submitted minus CRM received, in days. Blank when either date is missing.", "Claude")

R = "Quotes!$B$2:$B$%d" % QLAST
V = "Quotes!$J$2:$J$%d" % QLAST
A = "Quotes!$D$2:$D$%d" % QLAST
S = "Quotes!$K$2:$K$%d" % QLAST
VH = "Quotes!$I$2:$I$%d" % QLAST
DY = "Quotes!$P$2:$P$%d" % QLAST

# ---------- Open Requests ----------
os_ = wb.create_sheet("Open Requests")
oh = ["Rep", "Agency", "Buyer", "Request", "Received", "Due", "Status", "CRM RFQ #", "Next step"]
header(os_, 1, oh, [18, 24, 20, 40, 11, 11, 50, 13, 38])
for i, r in enumerate(OPEN, 2):
    z = i % 2 == 1
    for j, v in enumerate(r, 1):
        c = os_.cell(row=i, column=j, value=v)
        body(c, DT if j in (5, 6) else None, z, center if j in (5, 6, 8) else wrap)
    if "PAST DUE" in r[6]:
        os_.cell(row=i, column=7).font = Font(name=F, size=10, bold=True, color="C00000")
OLAST = len(OPEN) + 1
os_.freeze_panes = "B2"
os_.auto_filter.ref = f"A1:I{OLAST}"

# ---------- Data Flags ----------
fs = wb.create_sheet("Data Flags")
fh = ["#", "Type", "Issue", "Where", "Fix", "Owner"]
header(fs, 1, fh, [5, 12, 62, 34, 38, 16])
for i, r in enumerate(FLAGS, 2):
    z = i % 2 == 1
    for j, v in enumerate((i - 1,) + r, 1):
        body(fs.cell(row=i, column=j, value=v), None, z, center if j == 1 else wrap)
fs.freeze_panes = "A2"

# ---------- Method ----------
ms = wb.create_sheet("Method")
ms.column_dimensions["A"].width = 26
ms.column_dimensions["B"].width = 70
ms.column_dimensions["C"].width = 46
ms["A1"] = "How this report was built"; ms["A1"].font = title_f
rows = [
 ("Week", "Mon 9/21/2026 - Sun 9/27/2026 (no quotes dated on the weekend)"),
 ("Prepared", "9/28/2026 by Claude for Parker Marshall. Human review before sharing outside BBS."),
 ("Scope", "California state agencies only. Kim Quezon's 18 agencies and Aanya Samarasinghe's remaining state agencies, plus state work Parker handled."),
 ("Out of scope", "Body armor (BA-CA, BODYARMOR, BA-SRCWL), CSU / UC (Joey Magliato), federal, primes, local government, private."),
 ("What counts", "One row per request with a BBS quote built 9/21-9/25. Revisions and duplicate quote numbers counted once. Re-sent older quotes not counted."),
 ("Value", "Pre-tax subtotal. Priority: quote PDF in folder > Business Central quote total > bid-form inputs > CRM estimate."),
 ("Sources", "Folder tree (All RFQ's by Agency, file dates 9/21-9/27); Sales CRM bbs_rfqs (production, pulled 9/28); Business Central salesQuotes (documentDate 9/21-9/27, pulled 9/28); Parker's Sent Items (121 emails, 9/21-9/27)."),
 ("Blind spots", "No access to Kim's or Aanya's mailboxes, so 'sent' is taken from the CRM submitted date or the quote log. Quotes that exist only in someone's mailbox are missed."),
 ("Next week", "Re-run the same four sources for the new week and save as '<Monday date> Weekly Quote Report - CA State.xlsx' in this folder."),
]
for i, (k, v) in enumerate(rows, 3):
    ms.cell(row=i, column=1, value=k).font = bold_f
    c = ms.cell(row=i, column=2, value=v); c.font = b_f; c.alignment = wrap
    ms.cell(row=i, column=1).alignment = wrap
er = len(rows) + 4
ms.cell(row=er, column=1, value="Excluded on purpose").font = sec_f
header(ms, er + 1, ["Rep", "Item", "Reason"])
ms.column_dimensions["A"].width = 26
for i, r in enumerate(EXCLUDED, er + 2):
    for j, v in enumerate(r, 1):
        body(ms.cell(row=i, column=j, value=v), None, i % 2 == 1)

# ---------- Summary ----------
sm = wb.create_sheet("Summary", 0)
widths = [24, 12, 16, 14, 14, 11, 16, 14, 12]
for i, w in enumerate(widths, 1):
    sm.column_dimensions[get_column_letter(i)].width = w
sm["A1"] = "CA State Agencies - Weekly Quote Report"; sm["A1"].font = title_f
sm["A2"] = "Week of Mon 9/21/2026 - Sun 9/27/2026  |  Prepared 9/28/2026  |  State agencies only; values pre-tax"
sm["A2"].font = sub_f

sm["A4"] = "Rep scorecard"; sm["A4"].font = sec_f
header(sm, 5, ["Rep", "Quotes", "Quote value", "Average quote", "Largest quote", "On NASPO",
               "Value from quote PDFs", "Avg days received to submitted", "Open requests"])
sm.row_dimensions[5].height = 44
for i, rep in enumerate([KIM, AANYA, PARKER], 6):
    f = [rep,
         f"=COUNTIFS({R},A{i})",
         f"=SUMIFS({V},{R},A{i})",
         f"=IFERROR(C{i}/B{i},0)",
         f"=_xlfn.MAXIFS({V},{R},A{i})",
         f"=COUNTIFS({R},A{i},{VH},\"NASPO*\")",
         f"=SUMIFS({V},{R},A{i},{S},\"{DOC}\")",
         f"=IFERROR(AVERAGEIFS({DY},{R},A{i}),\"-\")",
         f"=COUNTIFS('Open Requests'!$A$2:$A${OLAST},A{i})"]
    for j, v in enumerate(f, 1):
        c = sm.cell(row=i, column=j, value=v)
        body(c, USD0 if j in (3, 4, 5, 7) else "0.0" if j == 8 else "0" if j in (2, 6, 9) else None,
             False, Alignment(horizontal="left" if j == 1 else "center", vertical="center"))
t = 9
tot = ["Total", "=SUM(B6:B8)", "=SUM(C6:C8)", "=IFERROR(C9/B9,0)", "=MAX(E6:E8)", "=SUM(F6:F8)",
       "=SUM(G6:G8)", f"=IFERROR(AVERAGE({DY}),\"-\")", "=SUM(I6:I8)"]
for j, v in enumerate(tot, 1):
    c = sm.cell(row=t, column=j, value=v)
    c.font, c.fill, c.border = bold_f, tot_fill, top2
    c.alignment = Alignment(horizontal="left" if j == 1 else "center")
    c.number_format = USD0 if j in (3, 4, 5, 7) else "0.0" if j == 8 else "0"
sm["A10"] = "Parker's CHP re-sends and the unsent Cal Fire Redding quote are tracked under Open Requests, not counted as completed."
sm["A10"].font = sub_f

# by agency
sm["A12"] = "By agency"; sm["A12"].font = sec_f
header(sm, 13, ["Agency", "Rep", "Quotes", "Quote value", "Share of value"])
ag_rows = [("CDCR", KIM), ("CMD", KIM), ("CHP", KIM), ("CDFW", KIM), ("SWRCB", KIM), ("CAL FIRE", KIM),
           ("DSH", AANYA), ("OEP", AANYA), ("DGS", AANYA), ("PDC", AANYA), ("DIR", AANYA)]
for i, (ag, rep) in enumerate(ag_rows, 14):
    f = [ag, rep.split()[0], f"=COUNTIFS({A},A{i})", f"=SUMIFS({V},{A},A{i})", f"=IFERROR(D{i}/$C$9,0)"]
    for j, v in enumerate(f, 1):
        c = sm.cell(row=i, column=j, value=v)
        body(c, "0" if j == 3 else USD0 if j == 4 else "0.0%" if j == 5 else None, i % 2 == 1,
             Alignment(horizontal="left" if j <= 2 else "center"))
al = 14 + len(ag_rows)
for j, v in enumerate(["Total", "", f"=SUM(C14:C{al-1})", f"=SUM(D14:D{al-1})", f"=IFERROR(D{al}/$C$9,0)"], 1):
    c = sm.cell(row=al, column=j, value=v)
    c.font, c.fill, c.border = bold_f, tot_fill, top2
    c.number_format = "0" if j == 3 else USD0 if j == 4 else "0.0%"
    c.alignment = Alignment(horizontal="left" if j <= 2 else "center")
sm.cell(row=al + 1, column=1, value=f'=IF(C{al}=B9,"Check: agency total ties to scorecard","CHECK: agency rows do not tie to scorecard - an agency is missing from this table")').font = sub_f

# value by source
vs = al + 3
sm.cell(row=vs, column=1, value="How solid the value is").font = sec_f
header(sm, vs + 1, ["Value source", "Confidence", "Quotes", "Quote value", "Share of value"])
for i, (src, cf) in enumerate([(DOC, "High"), (BC, "Medium"), (FORM, "Medium"), (CRM, "Low")], vs + 2):
    f = [src, cf, f"=COUNTIFS({S},A{i})", f"=SUMIFS({V},{S},A{i})", f"=IFERROR(D{i}/$C$9,0)"]
    for j, v in enumerate(f, 1):
        c = sm.cell(row=i, column=j, value=v)
        body(c, "0" if j == 3 else USD0 if j == 4 else "0.0%" if j == 5 else None, i % 2 == 1,
             Alignment(horizontal="left" if j <= 2 else "center"))

# watch list
wl = vs + 7
sm.cell(row=wl, column=1, value="Watch list").font = sec_f
watch = [
 "DSH Metro beds (SQ-104998) is the largest quote of the week and is not marked submitted in the CRM - confirm it went out.",
 "CalVet VHC Fresno RFQ was due 9/23; no quote found and the CRM record is still open.",
 "Likely win: BC SQ-105001 (SWRCB, Accepted 9/25) matches TaMara Skondras's 9/4 quote 26-170-005 - log the outcome.",
 "Cal Fire Redding turnouts: waiting on Brandon Self to say whether the buy must stay under $50K.",
 "Quote logs are nearly empty this week - see Data Flags.",
]
for i, w in enumerate(watch, wl + 1):
    sm.cell(row=i, column=1, value="-  " + w).font = b_f
sm.cell(row=wl + len(watch) + 2, column=1,
        value="Detail: Quotes tab (one row per quote)  |  Open Requests  |  Data Flags  |  Method (sources, scope, exclusions)").font = sub_f
sm.freeze_panes = "A6"

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

wb.save(OUT)
print("saved", OUT, "quotes", len(Q), "open", len(OPEN), "flags", len(FLAGS))
print("check total", round(sum(r[8] for r in Q), 2),
      {rep: round(sum(r[8] for r in Q if r[0] == rep), 2) for rep in (KIM, AANYA, PARKER)})
